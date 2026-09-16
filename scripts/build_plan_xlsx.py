"""Build plan/spain2027-plan.xlsx from the markdown plan files.

Run from the repo root: python3 scripts/build_plan_xlsx.py
Sheets: Deadlines, Monthly actions, Programs, Income log, Summary, How to use.
"""
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "plan" / "spain2027-plan.xlsx"
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
WORKSTREAMS = ["Programs", "Immigration", "Tax", "Income", "Trip", "Outreach",
               "Kindergarten", "Spanish", "Move", "Housing", "Work", "Dog"]
STATUSES = ["Not started", "In progress", "Waiting on", "At risk", "Overdue", "Done"]


def md_table(path):
    rows = []
    for line in path.read_text().splitlines():
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows[0], rows[1:]


def parse_due(text):
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", text)
    if m:
        return date(*map(int, m.groups())), ""
    m = re.match(r"(\d{4})-(\d{2})", text)
    if m:
        y, mo = map(int, m.groups())
        note = text[m.end():].strip(" ()")
        return date(y, mo, 1), ("Month only. " + note).strip()
    return None, text


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{ws.max_row}"


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def body_font(ws):
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_validations(ws, status_col, ws_col, last_row):
    dv_status = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv_ws = DataValidation(type="list", formula1='"' + ",".join(WORKSTREAMS) + '"', allow_blank=True)
    ws.add_data_validation(dv_status)
    ws.add_data_validation(dv_ws)
    dv_status.add(f"{status_col}2:{status_col}{last_row}")
    dv_ws.add(f"{ws_col}2:{ws_col}{last_row}")


def build_deadlines(wb):
    ws = wb.active
    ws.title = "Deadlines"
    header, rows = md_table(ROOT / "plan" / "deadlines.md")
    cols = ["Task", "Workstream", "Due", "Days left", "Status", "Type", "Confidence", "Source, checked", "Notes"]
    ws.append(cols)
    for r in rows:
        due_text, item, workstream, typ, conf, source, status = r[:7]
        due, note = parse_due(due_text)
        status_val = "Done" if status.lower().startswith("past") else ("Not started" if status == "Open" else status)
        ws.append([item, workstream, due, None, status_val, typ, conf, source, note])
    last = ws.max_row
    for i in range(2, last + 1):
        ws.cell(row=i, column=4).value = f'=IF(OR(C{i}="",E{i}="Done"),"",C{i}-TODAY())'
        ws.cell(row=i, column=3).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=5).fill = INPUT_FILL
        ws.cell(row=i, column=9).fill = INPUT_FILL
    style_header(ws, len(cols))
    body_font(ws)
    set_widths(ws, [70, 14, 12, 10, 13, 17, 20, 34, 40])
    add_validations(ws, "E", "B", last)
    rng = f"A2:I{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=['AND($D2<>"",$D2<0)'], fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['AND($D2<>"",$D2>=0,$D2<=14)'], fill=PatternFill("solid", fgColor="FFE699")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$E2="Done"'], font=Font(name=FONT, color="808080", strike=True)))
    return ws


def build_monthly(wb):
    ws = wb.create_sheet("Monthly actions")
    cols = ["Month", "Workstream", "Action", "Due", "Status", "Notes"]
    ws.append(cols)
    month = None
    workstream = None
    for line in (ROOT / "plan" / "monthly-plan.md").read_text().splitlines():
        m = re.match(r"## (\w+) (\d{4})", line)
        if m:
            month = datetime.strptime(f"{m.group(1)} {m.group(2)}", "%B %Y").date()
            continue
        if line and not line.startswith(("-", "#", " ")) and month is not None:
            workstream = line.strip()
            continue
        m = re.match(r"- \[( |x)\] (.*)", line)
        if m and month is not None:
            done = m.group(1) == "x"
            text = m.group(2)
            due = None
            d = re.search(r"Due (\d{2})-(\d{2})\.?", text)
            if d:
                due = date(month.year, int(d.group(1)), int(d.group(2)))
                text = text.replace(d.group(0), "").strip()
            ws.append([month, workstream, text, due, "Done" if done else "Not started", ""])
    last = ws.max_row
    for i in range(2, last + 1):
        ws.cell(row=i, column=1).number_format = "mmm yyyy"
        ws.cell(row=i, column=4).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=5).fill = INPUT_FILL
        ws.cell(row=i, column=6).fill = INPUT_FILL
    style_header(ws, len(cols))
    body_font(ws)
    set_widths(ws, [11, 14, 80, 12, 13, 40])
    add_validations(ws, "E", "B", last)
    rng = f"A2:F{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=['AND($D2<>"",$E2<>"Done",$D2<TODAY())'], fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['$E2="Done"'], font=Font(name=FONT, color="808080", strike=True)))
    return ws


def build_programs(wb):
    ws = wb.create_sheet("Programs")
    header, rows = md_table(ROOT / "plan" / "programs" / "tracker.md")
    ws.append(header)
    for r in rows:
        ws.append(r)
    style_header(ws, len(header))
    body_font(ws)
    set_widths(ws, [42, 24, 11, 40, 30, 50, 13])
    for i in range(2, ws.max_row + 1):
        ws.cell(row=i, column=4).fill = INPUT_FILL
        ws.cell(row=i, column=5).fill = INPUT_FILL
        ws.cell(row=i, column=6).fill = INPUT_FILL
    return ws


def build_income(wb):
    ws = wb.create_sheet("Income log")
    ws["A1"] = "Threshold EUR per month (2026 basis, family of three)"
    ws["B1"] = 4275
    ws["A2"] = "Exchange buffer"
    ws["B2"] = 0.15
    ws["A3"] = "Target USD per month"
    ws["B3"] = "=B1*(1+B2)*B6"
    ws["A4"] = "Source: chat 1, 2026-08-07. Threshold rises with the SMI each year; update B1 when the 2027 figure is published."
    ws["A6"] = "EUR to USD rate used for target"
    ws["B6"] = 1.144
    ws["A7"] = "Rate source: chat 1, 2026-08-07. Update to the current ECB rate."
    for r in (1, 2, 6):
        ws.cell(row=r, column=2).font = Font(name=FONT, color="0000FF")
        ws.cell(row=r, column=2).fill = INPUT_FILL
    ws["B2"].number_format = "0%"
    ws["B3"].number_format = "$#,##0"
    ws["B6"].number_format = "0.000"
    start = 9
    cols = ["Month", "Gross deposits USD", "EUR at month-end rate", "Rate used", "Above threshold", "Non-Spain share", "Notes"]
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(row=start, column=c, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    months = []
    y, m = 2026, 10
    for _ in range(15):
        months.append(date(y, m, 1))
        m += 1
        if m == 13:
            y, m = y + 1, 1
    for i, mo in enumerate(months):
        r = start + 1 + i
        ws.cell(row=r, column=1, value=mo).number_format = "mmm yyyy"
        ws.cell(row=r, column=2).number_format = "$#,##0"
        ws.cell(row=r, column=2).fill = INPUT_FILL
        ws.cell(row=r, column=4).fill = INPUT_FILL
        ws.cell(row=r, column=4).number_format = "0.000"
        ws.cell(row=r, column=3, value=f'=IF(OR(B{r}="",D{r}=""),"",B{r}/D{r})').number_format = "€#,##0"
        ws.cell(row=r, column=5, value=f'=IF(C{r}="","",IF(C{r}>=$B$1,"Yes","No"))')
        ws.cell(row=r, column=6).fill = INPUT_FILL
        ws.cell(row=r, column=6).number_format = "0%"
        ws.cell(row=r, column=7).fill = INPUT_FILL
    # example row
    r = start + 1
    ws.cell(row=r, column=2, value=6000)
    ws.cell(row=r, column=4, value=1.144)
    ws.cell(row=r, column=6, value=1.0)
    ws.cell(row=r, column=7, value="Example values. Replace with real October deposits.")
    ws.freeze_panes = f"A{start + 1}"
    for row in ws.iter_rows():
        for cell in row:
            if cell.font.color is None or cell.font.color.rgb in (None, "FF000000"):
                cell.font = Font(name=FONT, bold=cell.font.bold, color=cell.font.color)
    set_widths(ws, [30, 20, 22, 11, 16, 16, 50])
    return ws


def build_summary(wb, dl_last, ma_last):
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "Spain 2027 plan summary"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = "Today"
    ws["B2"] = "=TODAY()"
    ws["B2"].number_format = "yyyy-mm-dd"
    ws["A4"] = "Deadlines: open items by workstream"
    ws["A4"].font = Font(name=FONT, bold=True)
    hdr = ["Workstream", "Open", "Overdue", "Due in 14 days", "Done"]
    for c, h in enumerate(hdr, start=1):
        cell = ws.cell(row=5, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    r = 6
    for w in WORKSTREAMS:
        ws.cell(row=r, column=1, value=w)
        ws.cell(row=r, column=2, value=f'=COUNTIFS(Deadlines!$B$2:$B${dl_last},A{r},Deadlines!$E$2:$E${dl_last},"<>Done")')
        ws.cell(row=r, column=3, value=f'=COUNTIFS(Deadlines!$B$2:$B${dl_last},A{r},Deadlines!$E$2:$E${dl_last},"<>Done",Deadlines!$C$2:$C${dl_last},"<"&TODAY())')
        ws.cell(row=r, column=4, value=f'=COUNTIFS(Deadlines!$B$2:$B${dl_last},A{r},Deadlines!$E$2:$E${dl_last},"<>Done",Deadlines!$C$2:$C${dl_last},">="&TODAY(),Deadlines!$C$2:$C${dl_last},"<="&(TODAY()+14))')
        ws.cell(row=r, column=5, value=f'=COUNTIFS(Deadlines!$B$2:$B${dl_last},A{r},Deadlines!$E$2:$E${dl_last},"Done")')
        r += 1
    ws.cell(row=r, column=1, value="Total").font = Font(name=FONT, bold=True)
    for c in range(2, 6):
        col = get_column_letter(c)
        ws.cell(row=r, column=c, value=f"=SUM({col}6:{col}{r - 1})").font = Font(name=FONT, bold=True)
    r += 2
    ws.cell(row=r, column=1, value="Monthly actions: open items by month").font = Font(name=FONT, bold=True)
    r += 1
    for c, h in enumerate(["Month", "Open", "Done"], start=1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    r += 1
    y, m = 2026, 9
    for _ in range(14):
        d = date(y, m, 1)
        ws.cell(row=r, column=1, value=d).number_format = "mmm yyyy"
        ws.cell(row=r, column=2, value=f"=COUNTIFS('Monthly actions'!$A$2:$A${ma_last},A{r},'Monthly actions'!$E$2:$E${ma_last},\"<>Done\")")
        ws.cell(row=r, column=3, value=f"=COUNTIFS('Monthly actions'!$A$2:$A${ma_last},A{r},'Monthly actions'!$E$2:$E${ma_last},\"Done\")")
        r += 1
        m += 1
        if m == 13:
            y, m = y + 1, 1
    for row in ws.iter_rows():
        for cell in row:
            if cell.font.name != FONT:
                cell.font = Font(name=FONT, bold=cell.font.bold, size=cell.font.size, color=cell.font.color)
    set_widths(ws, [40, 10, 10, 16, 10])
    return ws


def build_howto(wb):
    ws = wb.create_sheet("How to use")
    lines = [
        "Spain 2027 relocation and program plan",
        "",
        "Sheets",
        "Summary: counts of open, overdue, and due-soon items. Formulas, do not edit.",
        "Deadlines: every dated item. Edit the yellow Status and Notes columns. Days left and highlighting update automatically: red is overdue, yellow is due within 14 days.",
        "Monthly actions: the action items for each month. Edit Status and Notes.",
        "Programs: the program tracker. Edit Status, Deadline, Missing items.",
        "Income log: enter each month's gross deposits and the rate used. The sheet shows whether the month clears the visa threshold. The first row is an example, replace it.",
        "",
        "Yellow cells are for you to edit. Everything else is either source data or a formula.",
        "Dates marked Month only had no exact day in the source; the day is set to the 1st.",
        "Confidence column: Confirmed means verified with the source on the date shown. Assumed means inferred and needs checking. Not verified means unknown.",
        "",
        "Regenerating: this file is built from the markdown plan files in the spain2027 repo by scripts/build_plan_xlsx.py. If you edit here and also want the repo updated, share the edited file back.",
        "Built " + date.today().isoformat(),
    ]
    for i, t in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=FONT, bold=(i == 1 or t in ("Sheets",)), size=14 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    return ws


def main():
    wb = Workbook()
    dl = build_deadlines(wb)
    ma = build_monthly(wb)
    build_programs(wb)
    build_income(wb)
    build_summary(wb, dl.max_row, ma.max_row)
    build_howto(wb)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
